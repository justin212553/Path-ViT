"""Redesigned PATH-ViT figures (PI feedback: no curves, horizontal flow, visual over text, show layers).
One inch-based spec -> two backends: an editable .pptx (import into Google Slides) and a PNG preview
drawn at true physical size so point sizes and overflow match the slide.
Page 50.4 x 28.35 in, the size of the existing Figure deck."""
import sys
from pathlib import Path

PAPER = Path(r"D:\wonse\Documents\Job\urban_datalab\PATH-ViT\paper")
W, H = 50.4, 28.35
DARK = "#202124"
WSI, WSI_T = "#4A86E8", "#E8F0FD"
RNA, RNA_T = "#9900FF", "#F3E6FF"
CLI, CLI_T = "#6AA84F", "#EDF6E9"
FOLD, FOLD_T, FOLD_TX = "#E69138", "#FCE8D2", "#7F3F00"   # one meaning only: fit on the training fold
GRAY_T = "#F1F3F4"
CLUSTER = ["#4E79A7", "#F28E2B", "#E15759", "#76B7B2", "#59A14F", "#EDC948",
           "#B07AA1", "#FF9DA7", "#9C755F", "#BAB0AC", "#2F4B7C"]


class Canvas:
    def __init__(self):
        self.items = []

    def box(self, x, y, w, h, text=None, fill=None, line=DARK, weight=6, size=48, bold=True, color=DARK,
            align="CENTER", kind="RECT"):
        self.items.append(dict(t="box", x=x, y=y, w=w, h=h, text=text, fill=fill, line=line, weight=weight,
                               size=size, bold=bold, color=color, align=align, kind=kind))

    def label(self, x, y, w, h, text, size=44, color=DARK, bold=True, align="CENTER"):
        self.box(x, y, w, h, text=text, fill=None, line=None, size=size, color=color, bold=bold, align=align)

    def arrow(self, x1, y1, x2, y2, color=DARK, weight=6, head="end"):
        self.items.append(dict(t="line", x1=x1, y1=y1, x2=x2, y2=y2, color=color, weight=weight, head=head))

    def seg(self, x1, y1, x2, y2, color=DARK, weight=6):
        self.arrow(x1, y1, x2, y2, color, weight, head=None)

    def image(self, path, x, y, w, h):
        self.items.append(dict(t="img", path=str(path), x=x, y=y, w=w, h=h))

    def badge(self, x, y, w, text="training fold"):
        self.box(x, y, w, 0.9, text=text, fill=FOLD_T, line=FOLD, weight=4, size=36, color=FOLD_TX)


def preprocess():
    p = Canvas()
    for name, col, tint, y, h in [("WSI", WSI, WSI_T, 0.8, 7.8), ("Molecular\nRNA-seq\n+ CNV", RNA, RNA_T, 9.2, 9.4),
                                  ("Clinico-\npathologic", CLI, CLI_T, 19.2, 7.0)]:
        p.box(0.8, y, 48.8, h, fill=tint, line=None)
        p.label(1.0, y + h / 2 - 1.6, 6.4, 3.2, name, size=72, color=col)
    # WSI band, center y 4.7
    p.image(PAPER / "tcga_paad_example_1.jpg", 8.2, 1.7, 4.4, 2.6)
    p.image(PAPER / "tcga_paad_example_2.jpg", 8.2, 4.5, 4.4, 2.6)
    p.arrow(13.0, 4.7, 14.8, 4.7)
    for path, (dx, dy) in zip(["r0003_c0031.jpg", "r0004_c0065.jpg", "r0006_c0027.jpg", "r0007_c0041.jpg"],
                              [(0, 0), (2.2, 0), (0, 2.2), (2.2, 2.2)]):
        p.image(PAPER / path, 15.2 + dx, 2.6 + dy, 2.0, 2.0)
    p.arrow(19.8, 4.7, 21.6, 4.7)
    p.box(22.0, 3.4, 5.0, 2.6, text="UNI2-h\n(frozen)", fill="#FFFFFF", line=WSI, size=54)
    p.arrow(27.4, 4.7, 29.2, 4.7)
    for k in range(5):
        p.box(29.6, 2.9 + k * 0.78, 5.0, 0.5, fill="#CFE0FA", line=WSI, weight=3)
    p.arrow(35.0, 4.7, 36.8, 4.7)
    p.badge(37.4, 1.3, 4.4)
    for k in range(11):
        r, c = divmod(k, 6)
        p.box(37.4 + c * 1.6 + (0.8 if r else 0), 2.9 + r * 1.75, 1.3, 1.3, fill=CLUSTER[k], line=DARK, weight=3)
    for x, w, t in [(8.2, 4.4, "Slides"), (15.2, 4.2, "Patches"), (22.0, 5.0, "Patch encoder"),
                    (29.6, 5.0, "Patch features"), (37.4, 9.3, "K = 11 clusters")]:
        p.label(x, 7.4, w, 0.9, t)
    # RNA-seq row, center y 12.2
    for i, v in enumerate([0.92, 0.85, 0.9, 0.35, 0.88, 0.8, 0.93, 0.6, 0.9, 0.86, 0.15, 0.89]):
        g = int(255 * v)
        p.box(8.2 + i * 0.55, 11.5, 0.55, 1.4, fill="#%02x%02x%02x" % (g, g, g), line=DARK, weight=2)
    p.label(8.2, 13.0, 6.6, 0.8, "1,500 genes", size=40)
    p.arrow(15.2, 12.2, 17.0, 12.2)
    for i, hgt in enumerate([2.4, 1.5, 0.9, 0.55, 0.35, 0.22, 0.14, 0.09]):
        p.box(17.4 + i * 0.5, 13.0 - hgt, 0.45, hgt, fill="#C9A0FF", line=RNA, weight=2)
    p.label(21.4, 10.8, 3.4, 0.9, "log2(x + 1)", size=40, color=RNA)
    p.arrow(21.8, 12.2, 24.4, 12.2)
    for i, hgt in enumerate([0.3, 0.8, 1.6, 2.4, 2.4, 1.6, 0.8, 0.3]):
        p.box(24.8 + i * 0.5, 13.0 - hgt, 0.45, hgt, fill="#C9A0FF", line=RNA, weight=2)
    p.badge(29.2, 9.9, 3.4)
    p.label(29.2, 10.9, 3.4, 0.9, "z-score", size=40, color=RNA)
    p.arrow(29.2, 12.2, 32.6, 12.2)
    for i, c in enumerate(["#2166AC", "#F7F7F7", "#B2182B", "#D1E5F0", "#FDDBC7", "#67A9CF", "#F4A582",
                           "#F7F7F7", "#4393C3", "#D6604D", "#F7F7F7", "#92C5DE"]):
        p.box(33.0 + i * 0.55, 11.5, 0.55, 1.4, fill=c, line=DARK, weight=2)
    p.label(33.0, 13.0, 6.6, 0.8, "standardized", size=40)
    p.arrow(39.8, 12.2, 41.6, 12.2)
    p.box(42.0, 11.1, 3.4, 2.2, text="RNA-seq\n1,500", fill="#FFFFFF", line=RNA, size=44)
    # CNV row, center y 16.2
    p.box(8.2, 15.6, 6.6, 1.2, fill="#FFFFFF", line=DARK, weight=4)
    for x, w, c in [(9.0, 0.9, "#B2182B"), (10.7, 0.5, "#2166AC"), (12.1, 1.2, "#B2182B"), (13.8, 0.4, "#2166AC")]:
        p.box(x, 15.6, w, 1.2, fill=c, line=None)
    p.label(8.2, 16.9, 6.6, 0.8, "copy-number gain / loss", size=36)
    p.badge(14.7, 13.9, 3.0)
    p.label(14.5, 14.8, 3.4, 0.8, "log2(c / 2)", size=36, color=RNA)
    p.arrow(15.2, 16.2, 17.0, 16.2)
    for i, c in enumerate(["#F4A582", "#F7F7F7", "#92C5DE", "#B2182B", "#F7F7F7", "#D1E5F0", "#FDDBC7", "#2166AC",
                           "#F7F7F7", "#D6604D", "#67A9CF", "#F7F7F7"]):
        p.box(17.4 + i * 0.55, 15.6, 0.55, 1.2, fill=c, line=DARK, weight=2)
    p.label(17.4, 16.9, 6.6, 0.8, "163 genes", size=40)
    p.label(23.9, 14.9, 4.0, 0.8, "category mean", size=36, color=RNA)
    p.arrow(24.4, 16.2, 27.4, 16.2)
    for i, c in enumerate(["#FDDBC7", "#D1E5F0", "#F7F7F7", "#B2182B", "#92C5DE", "#F4A582", "#F7F7F7", "#4393C3"]):
        p.box(27.8 + i * 0.55, 15.6, 0.55, 1.2, fill=c, line=DARK, weight=2)
    p.label(27.8, 16.9, 4.4, 0.8, "8 categories", size=40)
    p.arrow(32.6, 16.2, 41.6, 16.2)
    p.box(42.0, 15.1, 3.4, 2.2, text="CNV\n8", fill="#FFFFFF", line=RNA, size=44)
    p.arrow(45.4, 12.2, 46.0, 12.2)
    p.arrow(45.4, 16.2, 46.0, 16.2)
    p.box(46.0, 11.1, 3.2, 6.2, text="1,508-d\nmolecular\ninput", fill=RNA_T, line=RNA, size=40)
    # clinicopathologic band, center y 22.7
    for r, row in enumerate([["Age", "Sex", "T", "N"], ["M", "Grade", "Margin"]]):
        for c, t in enumerate(row):
            p.box(8.2 + c * 2.6, 20.0 + r * 1.4, 2.4, 1.1, text=t, fill="#FFFFFF", line=CLI, weight=4, size=40)
    for c, t in enumerate(["KRAS", "TP53", "SMAD4", "CDKN2A"]):
        p.box(8.2 + c * 2.6, 22.9, 2.4, 1.1, text=t, fill=GRAY_T, line=DARK, weight=3, size=34)
    p.label(8.2, 24.2, 10.2, 0.8, "mutation status, pancreatic cancer cohorts only", size=34)
    p.badge(19.3, 20.0, 3.2)
    p.label(18.7, 21.0, 4.4, 1.5, "z-score\n+ observed flag", size=36, color=CLI)
    p.arrow(19.0, 22.9, 22.6, 22.9)
    for i, v in enumerate(["#A1D99B", "#31A354", "#E5F5E0", "#74C476", "#FFFFFF", "#C7E9C0"]):
        p.box(23.0, 20.5 + i * 0.75, 1.6, 0.75, fill=v, line=DARK, weight=2)
        p.box(24.6, 20.5 + i * 0.75, 0.9, 0.75, fill=(DARK if i != 4 else "#FFFFFF"), line=DARK, weight=2)
    p.label(22.4, 25.1, 3.6, 0.8, "value | flag", size=34)
    p.arrow(25.9, 22.9, 28.6, 22.9)
    p.box(29.0, 21.6, 6.4, 2.6, text="Feature vector\nto linear term", fill="#FFFFFF", line=CLI, size=44)
    # footer
    p.box(0.8, 26.75, 3.2, 0.9, text="training fold", fill=FOLD_T, line=FOLD, weight=4, size=32, color=FOLD_TX)
    p.label(4.2, 26.65, 45.4, 1.1, "Statistics, imputation means, and cluster centroids are fit on the training patients "
            "of each fold, then applied unchanged to validation, test, and external patients.",
            size=38, bold=False, align="LEFT")
    return p


def architecture():
    p = Canvas()
    for name, col, cy in [("WSI", WSI, 5.0), ("Molecular\nRNA-seq\n+ CNV", RNA, 12.0), ("Clinico-\npathologic", CLI, 16.4)]:
        p.label(0.6, cy - 1.4, 5.6, 2.8, name, size=60, color=col)
    # WSI lane, y 5.0
    p.label(6.0, 1.5, 5.0, 0.9, "11 cluster tokens", size=40)
    for k in range(11):
        p.box(6.5, 2.7 + k * 0.42, 4.0, 0.32, fill=CLUSTER[k], line=None)
    p.label(6.5, 7.4, 4.0, 0.8, "1536-d each", size=36)
    p.arrow(10.9, 5.0, 12.5, 5.0)
    p.box(12.9, 3.8, 4.6, 2.4, text="Linear\n1536 \u2192 64", fill="#FFFFFF", line=WSI)
    p.arrow(17.5, 5.0, 19.1, 5.0)
    p.box(19.5, 3.8, 4.0, 2.4, text="LayerNorm", fill=GRAY_T, line=WSI)
    p.arrow(23.5, 5.0, 25.1, 5.0)
    for k in range(11):
        p.box(25.5, 2.7 + k * 0.42, 2.0, 0.32, fill=CLUSTER[k], line=None)
    p.label(25.0, 7.4, 3.0, 0.8, "11 \u00d7 64", size=36)
    p.label(27.9, 3.8, 3.6, 0.9, "Key, Value", size=40, color=WSI)
    p.arrow(27.9, 5.0, 31.5, 5.0, color=WSI)
    p.box(31.9, 3.0, 6.6, 4.0, text="Multi-head\ncross-attention", fill=WSI_T, line=DARK, weight=8, size=52)
    p.label(38.5, 3.4, 3.1, 0.9, "z_wsi, 64", size=36)
    p.arrow(38.5, 4.6, 41.6, 4.6)
    # RNA lane, y 12.0
    y0, h = 10.8, 2.4
    p.box(6.5, y0, 4.0, h, text="1,508-d\ninput", fill=RNA_T, line=RNA, size=44)
    prev = 10.5
    for x, w, t, f in [(11.3, 4.2, "Linear\n1508 \u2192 256", "#FFFFFF"), (16.3, 3.4, "LayerNorm", GRAY_T),
                       (20.5, 2.4, "GELU", GRAY_T), (23.7, 2.6, "Dropout\n0.3", GRAY_T),
                       (27.1, 4.0, "Linear\n256 \u2192 64", "#FFFFFF")]:
        p.arrow(prev, 12.0, x, 12.0)
        p.box(x, y0, w, h, text=t, fill=f, line=RNA, size=40)
        prev = x + w
    p.arrow(prev, 12.0, 31.9, 12.0)
    p.box(31.9, y0, 5.0, h, text="z_mol, 64", fill=RNA_T, line=RNA, size=48)
    p.arrow(34.4, 7.0, 34.4, 10.8, color=RNA, head="start")
    p.label(34.7, 8.4, 2.8, 0.9, "Query", size=40, color=RNA)
    p.seg(36.9, 12.0, 39.6, 12.0, color=RNA)
    p.seg(39.6, 5.6, 39.6, 12.0, color=RNA)
    p.arrow(39.6, 5.6, 41.6, 5.6, color=RNA)
    # risk head column
    p.box(41.6, 3.8, 7.2, 2.4, text="Concatenate, 128", fill="#FFFFFF", line=DARK)
    p.arrow(45.2, 6.2, 45.2, 7.2)
    p.box(41.6, 7.2, 7.2, 1.8, text="LayerNorm", fill=GRAY_T, line=DARK)
    p.arrow(45.2, 9.0, 45.2, 10.2)
    p.box(41.6, 10.2, 7.2, 2.4, text="Linear 128 \u2192 4", fill="#FFFFFF", line=DARK)
    p.arrow(45.2, 12.6, 45.2, 15.4)
    p.box(44.2, 15.4, 2.0, 2.0, text="+", fill="#FFFFFF", line=DARK, kind="OVAL", size=72)
    p.arrow(45.2, 17.4, 45.2, 19.0)
    p.box(41.6, 19.0, 7.2, 2.4, text="4 hazard logits", fill="#FFFFFF", line=DARK)
    p.arrow(45.2, 21.4, 45.2, 22.4)
    p.box(41.6, 22.4, 7.2, 2.4, text="NLL + Cox loss", fill=GRAY_T, line=DARK)
    # clinicopathologic lane, y 16.4
    p.box(6.5, 15.2, 4.0, 2.4, text="Feature\nvector", fill=CLI_T, line=CLI, size=44)
    p.arrow(10.5, 16.4, 11.3, 16.4)
    p.box(11.3, 15.2, 5.6, 2.4, text="Linear \u2192 1\nzero-initialized", fill="#FFFFFF", line=CLI, size=40)
    p.arrow(16.9, 16.4, 44.2, 16.4, color=CLI)
    p.label(28.0, 15.3, 10.0, 0.9, "added to every logit", size=36, color=CLI)
    # notes
    for i, t in enumerate(["M1, M5, and the -SA variants replace cross-attention with a query-free self-attention block.",
                           "M2, M3, and M6 have no WSI lane. M1, M3, and M5 have no molecular lane. M1, M2, and M4 have no clinicopathologic lane.",
                           "TCGA-BRCA uses UNI, version 1, at 1024-d, and an 882-gene RNA-seq input without CNV."]):
        p.label(6.5, 20.4 + i * 1.3, 33.0, 1.1, t, size=38, bold=False, align="LEFT")
    return p


# ----------------------------------------------------------------- backends
def to_pptx(canvases, out):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.oxml.ns import qn
    from lxml import etree

    def col(h):
        return RGBColor.from_string(h.lstrip("#").upper())

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    for cv in canvases:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        for it in cv.items:
            if it["t"] == "img":
                s.shapes.add_picture(it["path"], Inches(it["x"]), Inches(it["y"]), Inches(it["w"]), Inches(it["h"]))
            elif it["t"] == "box":
                kind = MSO_SHAPE.OVAL if it["kind"] == "OVAL" else MSO_SHAPE.RECTANGLE
                sh = s.shapes.add_shape(kind, Inches(it["x"]), Inches(it["y"]), Inches(it["w"]), Inches(it["h"]))
                sh.shadow.inherit = False
                if it["fill"]:
                    sh.fill.solid(); sh.fill.fore_color.rgb = col(it["fill"])
                else:
                    sh.fill.background()
                if it["line"]:
                    sh.line.color.rgb = col(it["line"]); sh.line.width = Pt(it["weight"])
                else:
                    sh.line.fill.background()
                tf = sh.text_frame
                tf.word_wrap = True
                tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                for m in ("margin_left", "margin_right"):
                    setattr(tf, m, Inches(0.12))
                tf.margin_top = tf.margin_bottom = Inches(0.05)
                if it["text"]:
                    for i, line in enumerate(it["text"].split("\n")):
                        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                        para.alignment = {"CENTER": PP_ALIGN.CENTER, "LEFT": PP_ALIGN.LEFT}[it["align"]]
                        r = para.add_run(); r.text = line
                        r.font.name = "Arial"; r.font.size = Pt(it["size"]); r.font.bold = it["bold"]
                        r.font.color.rgb = col(it["color"])
            else:
                x1, y1, x2, y2 = it["x1"], it["y1"], it["x2"], it["y2"]
                cn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
                cn.line.color.rgb = col(it["color"]); cn.line.width = Pt(it["weight"])
                if it["head"]:
                    ln = cn.line._get_or_add_ln()
                    tag = "a:tailEnd" if it["head"] == "end" else "a:headEnd"
                    el = etree.SubElement(ln, qn(tag)); el.set("type", "triangle"); el.set("w", "med"); el.set("len", "med")
    prs.save(out)


def to_png(cv, out, dpi=40):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import Rectangle, Ellipse
    import matplotlib.image as mpimg
    fig = plt.figure(figsize=(W, H), dpi=dpi)
    ax = fig.add_axes([0, 0, 1, 1]); ax.set_xlim(0, W); ax.set_ylim(H, 0); ax.axis("off")
    fig.patch.set_facecolor("white")
    for it in cv.items:
        if it["t"] == "img":
            ax.imshow(mpimg.imread(it["path"]), extent=(it["x"], it["x"] + it["w"], it["y"] + it["h"], it["y"]), zorder=2)
        elif it["t"] == "box":
            kw = dict(facecolor=it["fill"] or "none", edgecolor=it["line"] or "none",
                      linewidth=it["weight"] if it["line"] else 0)
            if it["kind"] == "OVAL":
                ax.add_patch(Ellipse((it["x"] + it["w"] / 2, it["y"] + it["h"] / 2), it["w"], it["h"], **kw))
            else:
                ax.add_patch(Rectangle((it["x"], it["y"]), it["w"], it["h"], zorder=1 if it["fill"] in (WSI_T, RNA_T, CLI_T) else 3, **kw))
            if it["text"]:
                if it["align"] == "LEFT":
                    tx, ha = it["x"] + 0.12, "left"
                else:
                    tx, ha = it["x"] + it["w"] / 2, "center"
                t = ax.text(tx, it["y"] + it["h"] / 2, it["text"], zorder=4, ha=ha, va="center", fontsize=it["size"],
                            fontweight="bold" if it["bold"] else "normal", color=it["color"], family="Arial",
                            linespacing=1.0)
                it["_artist"] = t
        else:
            style = "-|>" if it["head"] else "-"
            a, b = (it["x1"], it["y1"]), (it["x2"], it["y2"])
            if it["head"] == "start":
                a, b = b, a
            ax.annotate("", xy=b, xytext=a, arrowprops=dict(arrowstyle=style, color=it["color"],
                        lw=it["weight"], mutation_scale=60, shrinkA=0, shrinkB=0))
    fig.canvas.draw()
    # overflow check: text extent vs its box
    r = fig.canvas.get_renderer()
    warn = []
    for it in cv.items:
        if "_artist" in it:
            bb = it["_artist"].get_window_extent(r)
            wi, hi = bb.width / dpi, bb.height / dpi
            if wi > it["w"] - 0.1 or hi > it["h"] + 0.05:
                warn.append(f"  overflow {it['text']!r}: text {wi:.2f}x{hi:.2f} in, box {it['w']:.2f}x{it['h']:.2f}")
    fig.savefig(out, dpi=dpi)
    plt.close(fig)
    return warn


if __name__ == "__main__":
    out = Path(sys.argv[1])
    cvs = [preprocess(), architecture()]
    for name, cv in zip(["preprocess_v2", "architecture_v2"], cvs):
        w = to_png(cv, out / f"{name}.png")
        print(name, "overflow warnings:", len(w)); print("\n".join(w))
    to_pptx(cvs, out / "figures_v2.pptx")
    print("saved", out / "figures_v2.pptx")
